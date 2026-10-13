#!/usr/bin/env python3
"""
Multi-Format Document Parser & Ingestion Driver
Implementerer spesifikasjonen i tools/TOOL_SKILL_DOCUMENT_PARSER.md.
Støtter PDF (pypdf), Excel (.xlsx via openpyxl), CSV (auto delimiter/encoding) og Markdown.
Inkluderer paragraf-bevarende chunking (§) for RAG-oppslag.
"""

from __future__ import annotations

import csv
import io
import os
from pathlib import Path
import re
import sys
from typing import Any, Dict, List, Optional

# Sikre UTF-8 i konsoll på Windows
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def chunk_legal_text(text: str, source_filename: str) -> List[Dict[str, Any]]:
    """Deler juridiske tekster per § (Paragraf) for optimal RAG-kontekst."""
    pattern = r"(§\s*\d+[-\w]*\.?\s*[^§]+)"
    matches = re.findall(pattern, text)

    result = []
    if matches:
        for idx, chunk in enumerate(matches):
            sec_match = re.search(r"§\s*\d+[-\w]*", chunk)
            section_ref = sec_match.group(0) if sec_match else "Generelt"
            clean_chunk = chunk.strip()
            if len(clean_chunk) > 20:
                result.append({
                    "source_file": source_filename,
                    "chunk_id": f"{source_filename}#chunk_{idx+1}",
                    "section_ref": section_ref,
                    "content": clean_chunk,
                    "char_count": len(clean_chunk)
                })
    else:
        # Fallback: del opp per avsnitt (dobbel linjeskift)
        paragraphs = [p.strip() for p in text.split("\n\n") if len(p.strip()) > 30]
        for idx, p in enumerate(paragraphs):
            result.append({
                "source_file": source_filename,
                "chunk_id": f"{source_filename}#para_{idx+1}",
                "section_ref": "Generelt",
                "content": p,
                "char_count": len(p)
            })

    return result


def parse_csv_file(file_path: Path) -> List[Dict[str, Any]]:
    """Leser CSV med automatisk tegnsett-deteksjon (utf-8 vs latin1) og skilletegn (; vs ,)."""
    raw_bytes = file_path.read_bytes()
    for enc in ["utf-8", "windows-1252", "iso-8859-1"]:
        try:
            content = raw_bytes.decode(enc)
            # Finn skilletegn
            sample = content[:4096]
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            reader = csv.DictReader(io.StringIO(content), dialect=dialect)
            return list(reader)
        except Exception:
            continue
    return []


def parse_markdown_or_txt(file_path: Path) -> str:
    """Leser Markdown og ren tekst med UTF-8 fallback."""
    raw = file_path.read_bytes()
    for enc in ["utf-8", "utf-8-sig", "windows-1252", "iso-8859-1"]:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_pdf_file(file_path: Path) -> str:
    """Leser tekst fra PDF via pypdf dersom tilgjengelig."""
    try:
        import pypdf
        reader = pypdf.PdfReader(str(file_path))
        pages_text = []
        for p_idx, page in enumerate(reader.pages):
            txt = page.extract_text()
            if txt:
                pages_text.append(f"--- Side {p_idx+1} ---\n{txt}")
        return "\n\n".join(pages_text)
    except Exception as e:
        return f"[Feil under lesing av PDF {file_path.name}: {e}]"


def parse_excel_file(file_path: Path) -> Dict[str, List[Dict[str, Any]]]:
    """Leser Excel (.xlsx) per ark via openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        sheets_data = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(c or f"Col_{i}") for i, c in enumerate(rows[0])]
            records = []
            for r in rows[1:]:
                if any(r):
                    records.append(dict(zip(headers, r)))
            sheets_data[sheet_name] = records
        return sheets_data
    except Exception as e:
        return {"error": [{"message": str(e)}]}


def parse_document(file_path: str | Path) -> Dict[str, Any]:
    """Universell parsing for alle støttede formater."""
    p = Path(file_path)
    if not p.exists():
        raise FileNotFoundError(f"Fil ikke funnet: {p}")

    ext = p.suffix.lower()
    if ext in [".md", ".txt"]:
        raw_text = parse_markdown_or_txt(p)
        chunks = chunk_legal_text(raw_text, p.name)
        return {
            "format": "markdown_text",
            "filename": p.name,
            "raw_text_length": len(raw_text),
            "chunk_count": len(chunks),
            "chunks": chunks
        }
    elif ext == ".csv":
        rows = parse_csv_file(p)
        return {
            "format": "csv",
            "filename": p.name,
            "row_count": len(rows),
            "sample_row": rows[0] if rows else {}
        }
    elif ext in [".xlsx", ".xlsm"]:
        sheets = parse_excel_file(p)
        return {
            "format": "excel",
            "filename": p.name,
            "sheets": list(sheets.keys()),
            "total_records": sum(len(v) for v in sheets.values())
        }
    elif ext == ".pdf":
        raw_text = parse_pdf_file(p)
        chunks = chunk_legal_text(raw_text, p.name)
        return {
            "format": "pdf",
            "filename": p.name,
            "chunk_count": len(chunks),
            "chunks": chunks
        }
    else:
        return {
            "format": "unsupported",
            "filename": p.name,
            "error": f"Format {ext} støttes ikke direkte."
        }


if __name__ == "__main__":
    if len(sys.argv) > 1:
        target = Path(sys.argv[1])
        res = parse_document(target)
        print(f"Format: {res['format']}, Fil: {res['filename']}")
        if "chunk_count" in res:
            print(f"Genererte biter: {res['chunk_count']}")
    else:
        print("Bruk: python tools/document_parser.py <sti-til-fil>")
